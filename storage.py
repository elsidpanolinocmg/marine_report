# storage.py
import os
import json
from docx import Document
from docx.shared import Inches
import re

def save_session_json(path: str, session: dict):
    """Save the whole session (pdf_path, articles, annotations, meta) to JSON."""
    with open(path, "w", encoding="utf-8") as f:
        json.dump(session, f, indent=2, ensure_ascii=False)

def load_session_json(path: str):
    """Load session JSON and return the dict."""
    with open(path, "r", encoding="utf-8") as f:
        return json.load(f)

def save_article_docx(outpath: str, article: dict):
    doc = Document()
    
    table = doc.add_table(rows=5, cols=2)
    table.autofit = False
    for row in table.rows:
        row.cells[0].width = Inches(1.5)
        row.cells[1].width = Inches(4.5)

    labels = ["Article Title:", "Author:", "Date Published:", "Issue No:", "Article Body:"]
    values = [
        article.get("title", ""),
        article.get("author", ""),
        article.get("date", ""),
        article.get("issue", ""),
        "\n\n".join(article.get("bodies", []))
    ]

    for i in range(5):
        table.cell(i, 0).text = labels[i]
        table.cell(i, 1).text = values[i]

    # Add images after table
    if article.get("images"):
        doc.add_paragraph("\nImages:")
        for img in article["images"]:
            path = img.get("path")
            if path and os.path.exists(path):
                doc.add_picture(path, width=Inches(4))

    doc.save(outpath)

def safe_filename(s: str) -> str:
    s2 = (s or "untitled").strip()
    s2 = re.sub(r'[\\/*?:"<>|]', "", s2)
    return s2[:120] or "untitled"
