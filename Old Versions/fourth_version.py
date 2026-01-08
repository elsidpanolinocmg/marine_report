import os
import re
import json
import uuid
import pytesseract
from tkinter import *
from tkinter import filedialog, messagebox
from pdf2image import convert_from_path
from PIL import Image, ImageTk
from docx import Document
from docx.shared import Inches

# configure tesseract (adjust if installed elsewhere)
pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

OUTPUT_DIR = "Processed_Articles"
if not os.path.exists(OUTPUT_DIR):
    os.makedirs(OUTPUT_DIR)


def clean_ocr_text(text: str) -> str:
    """Fix broken OCR text: hyphens, newlines, spacing"""
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)  # join hyphen-split words
    text = re.sub(r"\n+", "\n", text)             # collapse multiple newlines
    text = re.sub(r"(?<!\n)\n(?!\n)", " ", text)  # single \n → space
    return text.strip()


class NewspaperOCRApp:
    def __init__(self, root):
        self.root = root
        self.root.title("Newspaper OCR App")
        self.root.geometry("1300x750")

        # state
        self.pages = []          # [{image, tk_image}]
        self.current_page = 0
        self.annotations = []    # [{article_id, page_index, bbox_orig, type, text}]
        self.articles = {}       # {article_id: {...}}
        self.current_article_id = None

        # ---- LEFT FRAME ----
        left = Frame(root)
        left.pack(side="left", fill="both", expand=True)

        self.canvas = Canvas(left, bg="gray")
        self.canvas.pack(fill="both", expand=True)

        nav = Frame(left)
        nav.pack(fill="x")
        Button(nav, text="Prev Page", command=self.prev_page).pack(side="left")
        Button(nav, text="Next Page", command=self.next_page).pack(side="left")
        self.page_label = Label(nav, text="Page 0/0")
        self.page_label.pack(side="left", padx=10)

        # zoom & pan
        self.scale = 1.0
        self.canvas.bind("<MouseWheel>", self.zoom)
        self.canvas.bind("<ButtonPress-1>", self.pan_start)
        self.canvas.bind("<B1-Motion>", self.pan_move)

        # draw selection (right mouse button)
        self.start_x = self.start_y = None
        self.rect = None
        self.canvas.bind("<ButtonPress-3>", self.on_rect_start)
        self.canvas.bind("<B3-Motion>", self.on_rect_draw)
        self.canvas.bind("<ButtonRelease-3>", self.on_rect_release)

        # ---- RIGHT FRAME ----
        right = Frame(root, width=400)
        right.pack(side="right", fill="y")

        Label(right, text="Articles").pack()
        self.article_listbox = Listbox(right)
        self.article_listbox.pack(fill="x", padx=5, pady=5)
        self.article_listbox.bind("<<ListboxSelect>>", self.on_article_select)

        Button(right, text="New Article", command=self.add_article).pack(fill="x", padx=5, pady=2)
        Button(right, text="Delete Article", command=self.delete_article).pack(fill="x", padx=5, pady=2)
        Button(right, text="Prev Article", command=self.prev_article).pack(fill="x", padx=5, pady=2)
        Button(right, text="Next Article", command=self.next_article).pack(fill="x", padx=5, pady=2)

        self.editor = Text(right, height=20, wrap="word")
        self.editor.pack(fill="both", expand=True, padx=5, pady=5)

        Button(right, text="Update Article", command=self.update_article).pack(fill="x", padx=5, pady=2)
        Button(right, text="Save All (Word)", command=self.save_all_word).pack(fill="x", padx=5, pady=2)
        Button(right, text="Export JSON", command=self.save_json).pack(fill="x", padx=5, pady=2)

        # menu
        menubar = Menu(root)
        filemenu = Menu(menubar, tearoff=0)
        filemenu.add_command(label="Open PDF", command=self.load_pdf)
        filemenu.add_separator()
        filemenu.add_command(label="Exit", command=root.quit)
        menubar.add_cascade(label="File", menu=filemenu)
        root.config(menu=menubar)

    # ================= PDF =================
    def load_pdf(self):
        filepath = filedialog.askopenfilename(filetypes=[("PDF files", "*.pdf")])
        if not filepath:
            return

        self.pages.clear()
        images = convert_from_path(filepath, dpi=150)
        for img in images:
            self.pages.append({"image": img, "tk_image": None})

        self.current_page = 0
        self.show_page()

    def show_page(self):
        if not self.pages:
            return
        page = self.pages[self.current_page]

        w, h = int(page["image"].width * self.scale), int(page["image"].height * self.scale)
        resized = page["image"].resize((w, h))
        page["tk_image"] = ImageTk.PhotoImage(resized)

        self.canvas.delete("all")
        self.canvas.create_image(0, 0, anchor="nw", image=page["tk_image"])

        # draw annotations for this page
        for ann in self.annotations:
            if ann["page_index"] == self.current_page:
                x1, y1, x2, y2 = [c * self.scale for c in ann["bbox_orig"]]
                color = {"Title": "blue", "Author": "orange", "Date": "purple", "Issue": "brown"}.get(ann["type"], "green")
                self.canvas.create_rectangle(x1, y1, x2, y2, outline=color, width=2)

        self.page_label.config(text=f"Page {self.current_page+1}/{len(self.pages)}")

    def next_page(self):
        if self.current_page < len(self.pages) - 1:
            self.current_page += 1
            self.show_page()

    def prev_page(self):
        if self.current_page > 0:
            self.current_page -= 1
            self.show_page()

    # ================= Canvas events =================
    def pan_start(self, event):
        self.canvas.scan_mark(event.x, event.y)

    def pan_move(self, event):
        self.canvas.scan_dragto(event.x, event.y, gain=1)

    def zoom(self, event):
        if event.delta > 0:
            self.scale *= 1.1
        else:
            self.scale /= 1.1
        self.show_page()

    def on_rect_start(self, event):
        self.start_x = self.canvas.canvasx(event.x)
        self.start_y = self.canvas.canvasy(event.y)
        self.rect = self.canvas.create_rectangle(
            self.start_x, self.start_y, self.start_x, self.start_y,
            outline="yellow", width=2
        )

    def on_rect_draw(self, event):
        cur_x = self.canvas.canvasx(event.x)
        cur_y = self.canvas.canvasy(event.y)
        self.canvas.coords(self.rect, self.start_x, self.start_y, cur_x, cur_y)

    def on_rect_release(self, event):
        x1 = self.start_x
        y1 = self.start_y
        x2 = self.canvas.canvasx(event.x)
        y2 = self.canvas.canvasy(event.y)

        if abs(x2 - x1) < 10 or abs(y2 - y1) < 10:
            return

        if not self.current_article_id:
            messagebox.showerror("Error", "Select or create an article first.")
            return

        bbox_orig = [c / self.scale for c in (x1, y1, x2, y2)]
        crop = self.pages[self.current_page]["image"].crop(bbox_orig)
        raw_text = pytesseract.image_to_string(crop, config="--psm 6")
        text = clean_ocr_text(raw_text)

        # popup type selection
        top = Toplevel(self.root)
        top.title("Select Type")
        Label(top, text="Choose annotation type:").pack()
        var = StringVar(value="Body")
        for t in ["Title", "Author", "Date", "Issue", "Body"]:
            Radiobutton(top, text=t, variable=var, value=t).pack(anchor="w")
        Button(top, text="OK", command=top.destroy).pack()
        top.wait_window()
        atype = var.get()

        ann = {
            "article_id": self.current_article_id,
            "page_index": self.current_page,
            "bbox_orig": bbox_orig,
            "type": atype,
            "text": text
        }
        self.annotations.append(ann)
        self.articles[self.current_article_id]["annotations"].append(ann)
        self.update_article_text()
        self.show_page()

    # ================= Articles =================
    def add_article(self):
        article_id = str(uuid.uuid4())
        self.articles[article_id] = {
            "title": f"Article {len(self.articles)+1}",
            "author": "",
            "date": "",
            "issue": "",
            "bodies": [],
            "annotations": []
        }
        self.refresh_article_list()
        self.article_listbox.selection_clear(0, END)
        self.article_listbox.selection_set(END)
        self.current_article_id = article_id
        self.load_article_into_editor(article_id)

    def delete_article(self):
        sel = self.article_listbox.curselection()
        if not sel:
            return
        idx = sel[0]
        aid = list(self.articles.keys())[idx]
        del self.articles[aid]
        self.annotations = [a for a in self.annotations if a["article_id"] != aid]
        self.refresh_article_list()
        self.editor.delete("1.0", END)
        self.current_article_id = None
        self.show_page()

    def on_article_select(self, event):
        sel = self.article_listbox.curselection()
        if not sel:
            return
        idx = sel[0]
        aid = list(self.articles.keys())[idx]
        self.current_article_id = aid
        self.load_article_into_editor(aid)

    def prev_article(self):
        sel = self.article_listbox.curselection()
        if not sel:
            return
        idx = sel[0]
        if idx > 0:
            self.article_listbox.selection_clear(0, END)
            self.article_listbox.selection_set(idx-1)
            self.on_article_select(None)

    def next_article(self):
        sel = self.article_listbox.curselection()
        if not sel:
            return
        idx = sel[0]
        if idx < self.article_listbox.size()-1:
            self.article_listbox.selection_clear(0, END)
            self.article_listbox.selection_set(idx+1)
            self.on_article_select(None)

    def load_article_into_editor(self, aid):
        art = self.articles[aid]
        self.editor.delete("1.0", END)
        txt = f"Title: {art['title']}\nAuthor: {art['author']}\nDate: {art['date']}\nIssue: {art['issue']}\n\n"
        bodies = "\n\n".join(art["bodies"])
        self.editor.insert("1.0", txt + bodies)

    def update_article_text(self):
        if not self.current_article_id:
            return
        art = self.articles[self.current_article_id]
        art["bodies"] = [a["text"] for a in art["annotations"] if a["type"] == "Body"]
        for a in art["annotations"]:
            if a["type"] == "Title":
                art["title"] = a["text"]
            if a["type"] == "Author":
                art["author"] = a["text"]
            if a["type"] == "Date":
                art["date"] = a["text"]
            if a["type"] == "Issue":
                art["issue"] = a["text"]
        self.load_article_into_editor(self.current_article_id)
        self.refresh_article_list()

    def refresh_article_list(self):
        self.article_listbox.delete(0, END)
        for aid, art in self.articles.items():
            self.article_listbox.insert(END, art["title"])

    def update_article(self):
        if not self.current_article_id:
            return
        txt = self.editor.get("1.0", END).strip()
        lines = txt.splitlines()
        art = self.articles[self.current_article_id]
        if lines and lines[0].startswith("Title:"):
            art["title"] = lines[0][6:].strip()
        if len(lines) > 1 and lines[1].startswith("Author:"):
            art["author"] = lines[1][7:].strip()
        if len(lines) > 2 and lines[2].startswith("Date:"):
            art["date"] = lines[2][5:].strip()
        if len(lines) > 3 and lines[3].startswith("Issue:"):
            art["issue"] = lines[3][6:].strip()
        body = "\n".join(lines[5:])
        art["bodies"] = [body]
        self.refresh_article_list()

    # ================= Save =================
    def save_all_word(self):
        for aid, art in self.articles.items():
            doc = Document()
            table = doc.add_table(rows=5, cols=2)
            table.autofit = False  # disable auto resize

            # set column widths
            for row in table.rows:
                row.cells[0].width = Inches(1.5)  # label column
                row.cells[1].width = Inches(5.0)  # value column

            labels = ["Article Title:", "Author:", "Date Published:", "Issue No:", "Article Body:"]
            values = [
                art.get("title", ""),
                art.get("author", ""),
                art.get("date", ""),
                art.get("issue", ""),
                "\n\n".join(art.get("bodies", []))
            ]

            for i in range(5):
                table.cell(i, 0).text = labels[i]
                table.cell(i, 1).text = values[i]

            safe_title = re.sub(r'[\\/*?:"<>|]', "", art["title"])[:50] or "Untitled"
            out = os.path.join(OUTPUT_DIR, f"{safe_title}.docx")
            doc.save(out)

        messagebox.showinfo("Saved", f"Articles saved to {OUTPUT_DIR}")

    def save_json(self):
        out = os.path.join(OUTPUT_DIR, "articles.json")
        with open(out, "w", encoding="utf-8") as f:
            json.dump({"articles": self.articles, "annotations": self.annotations}, f, indent=2, ensure_ascii=False)
        messagebox.showinfo("Saved", f"JSON saved to {out}")


if __name__ == "__main__":
    root = Tk()
    app = NewspaperOCRApp(root)
    root.mainloop()
